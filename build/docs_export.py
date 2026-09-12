#!/usr/bin/env python3
"""Turn the wiki mirror, the scene archive and the rule index into Word documents.

  python build/docs_export.py                      # writes to docs/ (gitignored)
  python build/docs_export.py --out "G:/My Drive/War of the Realms — Documents"
  python build/docs_export.py --only "Magic"       # sections whose name contains this

One .docx per wiki section (the section's hub page first, then every page under
it, alphabetical), plus "The Scene Archive" from scenes/ and "The Rule Index"
from the four index files. Each document opens with a title page and a
contents list; every page starts on a new page. Dropped into a Google Drive
folder they open in Google Docs as they are.

Only documents whose source text changed since the last run are rewritten
(docs/.manifest.json), so a Drive-synced folder is not re-uploaded every hour.
"""
import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for template tweaks)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
WIKI = ROOT / "wiki"
SCENES = ROOT / "scenes"

INLINE = re.compile(r"(\*\*.+?\*\*|\*(?!\s).+?(?<!\s)\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\[\[[^\]]+\]\])")

RULE_INDEX = [
    ("Natalie — Standing Rules", ROOT / "desktop" / "NATALIE.md"),
    ("Live Rules (resolved)", ROOT / "out" / "rules.live.full.md"),
    ("Docket (index view)", ROOT / "out" / "docket.md"),
    ("Conflicts", ROOT / "CONFLICTS.md"),
]
SCENE_SKIP = {"MANIFEST.md", "00_SUPERSEDED_wrong_names.md", "WOTR_AI_Writing_Tells_to_Avoid.md",
              "THE_KINGDOM_OF_KHARVEN_buri.md"}


# --------------------------------------------------------------------------
# markdown -> docx


def add_inline(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif part.startswith("[[") and part.endswith("]]"):
            par.add_run(part[2:-2]).italic = True
        elif part.startswith("[") and part.endswith(")"):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            par.add_run(m.group(1) if m else part)
        else:
            par.add_run(part)


def strip_frontmatter(md: str) -> tuple[str, dict]:
    md = md.replace("\r\n", "\n")
    meta = {}
    if md.startswith("---\n"):
        end = md.find("\n---\n", 4)
        if end != -1:
            for ln in md[4:end].split("\n"):
                if ":" in ln:
                    k, v = ln.split(":", 1)
                    try:
                        meta[k.strip()] = json.loads(v.strip())
                    except Exception:
                        meta[k.strip()] = v.strip()
            md = md[end + 5:]
    return md, meta


def render_md(doc, md: str, title: str, heading_offset: int = 0) -> None:
    """Render one page. Its own H1 (== title) is skipped; the caller wrote the heading."""
    lines = md.split("\n")
    para: list[str] = []
    i = 0
    first_h1_seen = False

    def flush():
        if para:
            p = doc.add_paragraph()
            add_inline(p, "\n".join(para))
            para.clear()

    while i < len(lines):
        s = lines[i].rstrip()
        st = s.strip()
        if not st:
            flush(); i += 1; continue
        if st.startswith("```"):
            flush()
            j = i + 1
            code = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code.append(lines[j]); j += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code)); r.font.name = "Consolas"; r.font.size = Pt(9)
            i = j + 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", st)
        if m:
            flush()
            level = len(m.group(1))
            text = m.group(2).strip().strip("*")
            if level == 1 and not first_h1_seen:
                first_h1_seen = True
                if text.lower() == title.lower():
                    i += 1; continue
            doc.add_heading(text, level=min(level + heading_offset, 6) if level > 1 or first_h1_seen else 1)
            i += 1; continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", st):
            flush()
            p = doc.add_paragraph("⁂"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if st.startswith(">"):
            flush()
            q = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q.append(lines[i].strip()[1:].strip()); i += 1
            p = doc.add_paragraph(style="Intense Quote") if "Intense Quote" in [s.name for s in doc.styles] else doc.add_paragraph()
            add_inline(p, "\n".join(x for x in q))
            continue
        if st.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{2,}", lines[i + 1].strip()):
            flush()
            rows = []
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                if not re.match(r"^\|?\s*:?-{2,}", lines[j].strip()):
                    rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            if rows:
                width = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=width)
                t.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else "Table Grid"
                for ri, r in enumerate(rows):
                    for ci in range(width):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        add_inline(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            i = j; continue
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", s)
        if m:
            flush()
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if depth == 0 else "List Bullet 2")
            add_inline(p, m.group(2))
            i += 1; continue
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", s)
        if m:
            flush()
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Number" if depth == 0 else "List Number 2")
            add_inline(p, m.group(2))
            i += 1; continue
        para.append(st)
        i += 1
    flush()


# --------------------------------------------------------------------------
# document assembly


def new_document(title: str, subtitle: str, pages: list[str]) -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(11)
    for lvl, size in ((1, 20), (2, 15), (3, 12.5)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = "Georgia"
        h.font.size = Pt(size)
        h.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)

    # title page
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WAR OF THE REALMS"); r.font.size = Pt(13); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.font.size = Pt(30); r.bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle); r.font.size = Pt(12); r.italic = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Generated {date.today().isoformat()} from the Notion wiki. Edit there, not here.")
    r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # contents
    doc.add_heading("Contents", level=1)
    for n, t in enumerate(pages, 1):
        doc.add_paragraph(f"{n}.  {t}")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return doc


def add_page(doc, title: str, md: str, page_break: bool = True) -> None:
    doc.add_heading(title, level=1)
    render_md(doc, md, title)
    if page_break:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def safe_name(s: str) -> str:
    return re.sub(r"[\\/:*?\"<>|]+", "", s).strip()[:120]


def collect_wiki() -> dict[str, list[tuple[str, str]]]:
    """section -> [(title, markdown)], hub page first then alphabetical."""
    manifest = json.loads((WIKI / ".manifest.json").read_text(encoding="utf-8"))
    pages = []
    for v in manifest.values():
        p = WIKI / v["rel"]
        if not p.exists():
            continue
        md, meta = strip_frontmatter(p.read_text(encoding="utf-8"))
        pages.append({"title": v["title"], "section": v["rel"].split("/")[0], "md": md})
    folders = {p["section"] for p in pages}
    by_title = {p["title"]: p for p in pages}
    docs: dict[str, list] = defaultdict(list)
    for p in pages:
        if p["title"] in folders:
            continue  # it is a hub: rendered at the top of its own section's document
        docs[p["section"]].append((p["title"], p["md"]))
    out = {}
    for section, items in docs.items():
        items.sort(key=lambda x: x[0].lower())
        hub = by_title.get(section)
        if hub:
            items.insert(0, (section, hub["md"]))
        out[section] = items
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "docs"))
    ap.add_argument("--private-out", help="folder for The Rule Index document (keep it out of anything shared)")
    ap.add_argument("--only")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    man_path = out_dir / ".manifest.json"
    manifest = json.loads(man_path.read_text(encoding="utf-8")) if man_path.exists() else {}

    jobs: dict[str, tuple[str, list[tuple[str, str]]]] = {}
    for section, items in collect_wiki().items():
        jobs[section] = (f"{len(items)} pages from the wiki section", items)

    scenes = []
    # reading order from scenes/ARCS.md when it exists; anything unlisted follows alphabetically
    order = []
    arcs = SCENES / "ARCS.md"
    if arcs.exists():
        order = re.findall(r"(?m)^-\s+(\S+\.md)", arcs.read_text(encoding="utf-8"))
    files = sorted(SCENES.glob("*.md"), key=lambda p: (order.index(p.name) if p.name in order else len(order), p.name))
    for p in files:
        if p.name in SCENE_SKIP or p.name in ("CAST.md", "ARCS.md"):
            continue
        md, _ = strip_frontmatter(p.read_text(encoding="utf-8", errors="replace"))
        m = re.search(r"(?m)^#\s+(.+?)\s*$", md)
        scenes.append((m.group(1).strip().strip("*") if m else p.stem.replace("_", " "), md))
    if scenes:
        jobs["The Scene Archive"] = (f"{len(scenes)} scenes", scenes)

    idx = [(t, strip_frontmatter(p.read_text(encoding="utf-8"))[0]) for t, p in RULE_INDEX if p.exists()]
    if idx:
        jobs["The Rule Index"] = ("the standing rules, the live craft law, the docket and the conflicts", idx)

    if args.only:
        jobs = {k: v for k, v in jobs.items() if args.only.lower() in k.lower()}

    private_dir = Path(args.private_out) if args.private_out else out_dir
    private_dir.mkdir(parents=True, exist_ok=True)
    if args.private_out and (out_dir / "The Rule Index.docx").exists():
        (out_dir / "The Rule Index.docx").unlink()  # it lives in the private folder now

    written = 0
    for section, (subtitle, items) in sorted(jobs.items()):
        h = hashlib.sha256("\n".join(t + m for t, m in items).encode("utf-8")).hexdigest()[:16]
        fname = f"{safe_name(section)}.docx"
        dest = private_dir if section == "The Rule Index" else out_dir
        if not args.force and manifest.get(fname) == h and (dest / fname).exists():
            continue
        doc = new_document(section, subtitle, [t for t, _ in items])
        for n, (t, md) in enumerate(items):
            add_page(doc, t, md, page_break=n < len(items) - 1)
        doc.save(dest / fname)
        manifest[fname] = h
        written += 1
        print(f"  wrote {fname} ({len(items)} pages)")
    man_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"{written} document(s) written, {len(jobs) - written} unchanged -> {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
